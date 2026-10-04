import io
from datetime import datetime, timedelta
from fastapi import APIRouter, Depends, Response
from fastapi.responses import StreamingResponse
from sqlalchemy import select, func, and_
from sqlalchemy.orm import selectinload
from sqlalchemy.ext.asyncio import AsyncSession
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

from app.database import get_db
from app.models import Booking, BookingStatus, PaymentStatus, MeetingRoom, User
from app.deps import get_current_approver

router = APIRouter(prefix="/reports", tags=["Reports & Analytics"])

@router.get("/dashboard-stats")
async def get_dashboard_stats(
    db: AsyncSession = Depends(get_db),
    approver: User = Depends(get_current_approver)
):
    now = datetime.utcnow()
    start_of_today = datetime(now.year, now.month, now.day)
    end_of_today = start_of_today + timedelta(days=1)
    start_of_month = datetime(now.year, now.month, 1)

    # 1. Total bookings today
    today_stmt = select(func.count(Booking.id)).where(
        and_(
            Booking.start_time >= start_of_today,
            Booking.start_time < end_of_today
        )
    )
    today_res = await db.execute(today_stmt)
    bookings_today = today_res.scalar() or 0

    # 2. Total revenue this month
    rev_stmt = select(func.sum(Booking.total_price)).where(
        and_(
            Booking.created_at >= start_of_month,
            Booking.payment_status == PaymentStatus.PAID
        )
    )
    rev_res = await db.execute(rev_stmt)
    monthly_revenue = float(rev_res.scalar() or 0.0)

    # 3. Pending approvals count
    pending_stmt = select(func.count(Booking.id)).where(Booking.status == BookingStatus.PENDING_APPROVAL)
    pending_res = await db.execute(pending_stmt)
    pending_approvals = pending_res.scalar() or 0

    # 4. Total rooms & occupancy rate
    total_rooms_stmt = select(func.count(MeetingRoom.id))
    total_rooms_res = await db.execute(total_rooms_stmt)
    total_rooms = total_rooms_res.scalar() or 1

    # Active bookings count right now
    active_now_stmt = select(func.count(Booking.id)).where(
        and_(
            Booking.status.in_([BookingStatus.APPROVED, BookingStatus.CHECKED_IN]),
            Booking.start_time <= now,
            Booking.end_time >= now
        )
    )
    active_now_res = await db.execute(active_now_stmt)
    active_now = active_now_res.scalar() or 0

    occupancy_rate = round((active_now / total_rooms) * 100, 1) if total_rooms > 0 else 0.0

    return {
        "bookings_today": bookings_today,
        "monthly_revenue": monthly_revenue,
        "pending_approvals": pending_approvals,
        "total_rooms": total_rooms,
        "occupied_rooms_now": active_now,
        "occupancy_rate_percent": occupancy_rate
    }

@router.get("/export/excel")
async def export_excel(
    db: AsyncSession = Depends(get_db),
    approver: User = Depends(get_current_approver)
):
    stmt = select(Booking).options(
        selectinload(Booking.room),
        selectinload(Booking.user)
    ).order_by(Booking.created_at.desc())
    res = await db.execute(stmt)
    bookings = res.scalars().all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Roomie Bookings Report"

    # Header Row
    headers = [
        "Booking Code", "Mode", "User/Guest", "Room Name",
        "Subject", "Start Time", "End Time", "Room Price",
        "Equipment Price", "Discount", "Total Price", "Booking Status", "Payment Status"
    ]
    ws.append(headers)

    # Styling header
    header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for b in bookings:
        user_name = b.user.full_name if b.user else f"{b.guest_name} (Guest)"
        ws.append([
            b.booking_code,
            b.booking_mode.value,
            user_name,
            b.room.name if b.room else "",
            b.subject,
            b.start_time.strftime("%Y-%m-%d %H:%M"),
            b.end_time.strftime("%Y-%m-%d %H:%M"),
            float(b.room_price),
            float(b.equipment_price),
            float(b.discount_amount),
            float(b.total_price),
            b.status.value,
            b.payment_status.value
        ])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    headers_response = {
        "Content-Disposition": "attachment; filename=roomie_bookings_report.xlsx"
    }
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers_response
    )

@router.get("/export/pdf")
async def export_pdf(
    db: AsyncSession = Depends(get_db),
    approver: User = Depends(get_current_approver)
):
    stmt = select(Booking).options(
        selectinload(Booking.room),
        selectinload(Booking.user)
    ).order_by(Booking.created_at.desc())
    res = await db.execute(stmt)
    bookings = res.scalars().all()

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
    story = []

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=18,
        spaceAfter=12,
        textColor=colors.HexColor('#1F2937')
    )

    story.append(Paragraph("Roomie System - Bookings Summary Report", title_style))
    story.append(Paragraph(f"Generated at: {datetime.utcnow().strftime('%Y-%m-%d %H:%M:%S')} UTC", styles['Normal']))
    story.append(Spacer(1, 15))

    # Table data
    data = [["Code", "Mode", "User/Guest", "Room", "Start Time", "Total (THB)", "Status"]]
    for b in bookings:
        u_name = b.user.full_name if b.user else f"{b.guest_name}"
        data.append([
            b.booking_code,
            b.booking_mode.value,
            u_name[:15],
            (b.room.name if b.room else "")[:15],
            b.start_time.strftime("%Y-%m-%d %H:%M"),
            f"{float(b.total_price):,.2f}",
            b.status.value
        ])

    t = Table(data, colWidths=[65, 55, 95, 95, 95, 65, 80])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E3A8A')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 9),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
        ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#F3F4F6')),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#D1D5DB')),
        ('FONTSIZE', (0, 1), (-1, -1), 8),
    ]))

    story.append(t)
    doc.build(story)
    buffer.seek(0)

    headers_response = {
        "Content-Disposition": "attachment; filename=roomie_bookings_report.pdf"
    }
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers=headers_response
    )
